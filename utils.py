"""
Utility and Export Helper Module for IT Ticket Tracking System
Handles Excel/CSV exports, date formatting, and styling utilities.
"""

import io
from datetime import datetime
from typing import List, Tuple
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December"
]

def export_to_csv(df: pd.DataFrame) -> bytes:
    """Convert dataframe to CSV bytes for download."""
    if df.empty:
        return b""
    return df.to_csv(index=False).encode('utf-8')

def export_to_excel(df: pd.DataFrame) -> bytes:
    """
    Export dataframe to a styled Excel (.xlsx) file in memory.
    Applies custom headers, column widths, and line wrapping.
    """
    output = io.BytesIO()
    
    if df.empty:
        # Create an empty excel file if df is empty
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tickets"
        ws.append(list(df.columns) if not df.columns.empty else ["No Data"])
        wb.save(output)
        return output.getvalue()
    
    # Remove internal ID columns if present for clean export
    export_df = df.copy()
    if 'Employee ID' in export_df.columns:
        export_df = export_df.drop(columns=['Employee ID'])

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        export_df.to_excel(writer, index=False, sheet_name="Tickets")
        
        workbook = writer.book
        worksheet = writer.sheets["Tickets"]
        
        # Define styles
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Navy Blue
        cell_font = Font(name="Segoe UI", size=10, color="000000")
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        
        # Style Header Row
        for col_idx, col_name in enumerate(export_df.columns, start=1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # Preferred Column Widths & Alignments
        col_width_map = {
            "Ticket ID": 18,
            "Date": 13,
            "Time": 12,
            "Who Was Helped": 22,
            "Issue": 45,
            "Employee": 20,
            "How It Was Resolved": 45,
            "Created At": 20,
            "Updated At": 20
        }
        
        # Format Data Rows
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column), start=2):
            for col_idx, cell in enumerate(row, start=1):
                cell.font = cell_font
                cell.border = thin_border
                col_name = str(export_df.columns[col_idx - 1])
                
                # Text alignment & wrapping
                if col_name in ["Issue", "How It Was Resolved"]:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                elif col_name in ["Date", "Time", "Ticket ID"]:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="top")
        
        # Set column widths
        for col_idx, col_name in enumerate(export_df.columns, start=1):
            col_letter = get_column_letter(col_idx)
            width = col_width_map.get(str(col_name), 20)
            worksheet.column_dimensions[col_letter].width = width

    return output.getvalue()

def get_year_range(df: pd.DataFrame) -> List[int]:
    """Extract list of unique years from ticket dataframe, including current year."""
    current_year = datetime.now().year
    years = {current_year}
    if not df.empty and 'Date' in df.columns:
        dates = pd.to_datetime(df['Date'], errors='coerce')
        valid_years = dates.dt.year.dropna().astype(int).unique()
        years.update(valid_years)
    return sorted(list(years), reverse=True)
